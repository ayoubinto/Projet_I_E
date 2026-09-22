import csv
import unicodedata
from datetime import date, datetime
from pathlib import Path

import pymupdf
from ollama import chat
from openpyxl import load_workbook
from paddleocr import PaddleOCR
from pydantic import BaseModel, field_validator


class Operation(BaseModel):
    reference_operation: str | None = None
    type_operation: str | None = None
    produit: str | None = None
    quantite: int | None = None
    prix_unitaire: float | None = None
    devise: str | None = None
    statut: str | None = None
    pays_origine: str | None = None
    pays_destination: str | None = None
    date_operation: str | None = None
    date_livraison: str | None = None


    @field_validator("type_operation")
    @classmethod
    def validate_type_operation(cls,value):
        if value is None:
            return None
        value = value.upper().strip()

        if value not in ["IMPORT","EXPORT"]:
            raise ValueError("type_operation doit être IMPORT ou EXPORT")

        return value

    @field_validator("statut")
    @classmethod
    def validate_statut(cls, value):
        if value is None:
            return None
        value = value.replace(" ", "_").replace("-", "_")
        statuts_valides = [
            "PLANIFIEE",
            "EN_TRANSIT",
            "LIVREE",
            "ANNULEE",
        ]

        value = value.upper().strip()

        value = "".join(
            char
            for char in unicodedata.normalize("NFD",value)
            if unicodedata.category(char) != "Mn"
        )

        if value not in statuts_valides:
            raise ValueError("Statut non valide")

        return value

    @field_validator("quantite")
    @classmethod
    def validate_quantite(cls,value):
        if value is not None and value <= 0:
            raise ValueError("La quantité doit être supérieure à 0")

        return value

    @field_validator("prix_unitaire")
    @classmethod
    def validate_prix(cls,value):
        if value is not None and value < 0:
            raise ValueError("Le prix unitaire ne peut pas être négatif")

        return value

    @field_validator("date_operation", "date_livraison")
    @classmethod
    def validate_date(cls, value):
        if value is None:
            return None

        datetime.strptime(value, "%d/%m/%Y")

        return value


# Chargé une seule fois
"""ocr = PaddleOCR(
    text_detection_model_name="PP-OCRv6_tiny_det",
    text_recognition_model_name="PP-OCRv6_tiny_rec",
    use_doc_orientation_classify=False,
    use_doc_unwarping=False,
    use_textline_orientation=False,
    engine="paddle"
)"""
ocr = None


def get_ocr():
    global ocr

    if ocr is None:
        print("Chargement de PaddleOCR...")

        ocr = PaddleOCR(
            text_detection_model_name="PP-OCRv6_tiny_det",
            text_recognition_model_name="PP-OCRv6_tiny_rec",
            use_doc_orientation_classify=False,
            use_doc_unwarping=False,
            use_textline_orientation=False,
            engine="paddle"
        )

    return ocr

def extract_text(file_path):
    extension = Path(file_path).suffix.lower()

    if extension == ".csv":
        print("CSV > module csv")
        return extract_csv_text(file_path
                                )
    if extension == ".xlsx":
        print("Excel > OpenPyXL")
        return extract_excel_text(file_path)

    if extension == ".pdf":
        native_text = extract_pdf_native_text(file_path)

        useful_text = "".join(native_text.split())

        if len(useful_text) >= 100:
            print("PDF numérique > PyMuPDF")
            return native_text

        print("PDF scanné > conversion en image > PaddleOCR")
        return extract_scanned_pdf_text(file_path)

    if extension in {".png", ".jpg", ".jpeg", ".webp"}:
        print("Image > PaddleOCR")
        return extract_ocr_text(file_path)

    raise ValueError(
        f"Format de fichier non supporté : {extension}"
    )

def extract_ocr_text(file_path):
    ocr_model = get_ocr()
    results = ocr_model.predict(file_path)

    texts = []

    for result in results:
        texts.extend(result["rec_texts"])

    return "\n".join(texts)

def extract_operation(file_path):

    # 1. OCR
    texte_extrait  = extract_text(file_path)
    print("\n===== TEXTE EXTRAIT =====")
    print(texte_extrait )

    # 2. LLM
    response = chat(
        model="llama3.2:1b",
        messages=[
            {
                "role": "system",
                "content": """
    Tu es un système d'extraction de données.
    
    Ta seule source de vérité est le texte fourni par l'utilisateur.
    
    N'utilise aucune valeur provenant d'exemples précédents.
    N'invente aucune information.
    Ne corrige pas les données selon ta logique.
    Ne déduis pas le pays d'origine ou de destination à partir du type IMPORT/EXPORT.
    
    Extrais exactement les valeurs présentes dans le document.
    
    Si une valeur est réellement absente, retourne null.
    """
            },
            {
                "role": "user",
                "content": f"""
    Extrais les champs suivants uniquement à partir du texte ci-dessous :
    
    - reference_operation
    - type_operation
    - produit
    - quantite
    - prix_unitaire
    - devise
    - statut
    - pays_origine
    - pays_destination
    - date_operation
    - date_livraison
    
    IMPORTANT :
    - copie les valeurs du document ;
    - ne remplace jamais une valeur présente ;
    - le libellé et sa valeur peuvent être séparés par | ou par un retour à la ligne.
    
    DOCUMENT :

    {texte_extrait }
    """
            }
        ],
        format=Operation.model_json_schema(),
        keep_alive="30m",
        options={
            "temperature": 0,
            "num_predict" : 256,
            "num_ctx" : 4096,
        }
    )
    print("\n===== PERFORMANCE OLLAMA =====")

    print(
        "Chargement modèle :",
        round((response.load_duration or 0) / 1_000_000_000, 2),
        "s"
    )

    print(
        "Analyse prompt :",
        round((response.prompt_eval_duration or 0) / 1_000_000_000, 2),
        "s"
    )

    print(
        "Génération réponse :",
        round((response.eval_duration or 0) / 1_000_000_000, 2),
        "s"
    )

    print("Tokens prompt :", response.prompt_eval_count)
    print("Tokens générés :", response.eval_count)

    # 3. Validation
    operation = Operation.model_validate_json(
        response.message.content
    )

    # 4. Retour Python
    return operation.model_dump()

def extract_pdf_native_text(pdf_path):
    document = pymupdf.open(pdf_path)

    pages_text = []

    for page in document:
        text = page.get_text("text", sort=True).strip()

        if text:
            pages_text.append(text)

    document.close()

    return "\n".join(pages_text)

def extract_scanned_pdf_text(pdf_path):
    document = pymupdf.open(pdf_path)
    texts = []

    try:
        for page_number, page in enumerate(document):
            print(f"OCR page PDF : {page_number + 1}")

            # Transformation de la page PDF en image
            pix = page.get_pixmap(matrix=pymupdf.Matrix(2, 2))

            temp_image = Path(pdf_path).with_name(
                f"_temp_page_{page_number + 1}.png"
            )

            pix.save(str(temp_image))

            try:
                page_text = extract_ocr_text(str(temp_image))

                if page_text.strip():
                    texts.append(page_text)

            finally:
                temp_image.unlink(missing_ok=True)

    finally:
        document.close()

    return "\n".join(texts)

def extract_excel_text(file_path):
    workbook = load_workbook(
        file_path,
        data_only=True
    )

    lines = []

    for sheet in workbook.worksheets:

        for row in sheet.iter_rows(values_only=True):

            values = []

            for value in row:

                if value is None:
                    continue

                # Date Excel
                if isinstance(value, (datetime,date)):
                    value = value.strftime("%d/%m/%Y")

                values.append(str(value).strip())

            if values:
                lines.append(" | ".join(values))

    workbook.close()

    return "\n".join(lines)

def extract_csv_text(file_path):
    with open(
        file_path,
        "r",
        encoding="utf-8-sig",
        newline=""
    ) as file:

        # Lire un petit échantillon pour détecter le séparateur
        sample = file.read(4096)
        file.seek(0)

        try:
            dialect = csv.Sniffer().sniff(
                sample,
                delimiters=",;\t|"
            )
        except csv.Error:
            dialect = csv.excel

        reader = csv.reader(file, dialect)

        lines = []

        for row in reader:
            values = [
                str(value).strip()
                for value in row
                if str(value).strip()
            ]

            if values:
                lines.append(" | ".join(values))

    return "\n".join(lines)


if __name__ == "__main__":

    file_path = "facture_csv.csv"

    data = extract_operation(file_path)

    print("\n===== RÉSULTAT FINAL =====")
    print(data)